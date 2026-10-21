"""The expense report as Excel and CSV files, in Korean or English."""

import csv
import io
import re
from datetime import date
from typing import Literal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import normalize as nz
from .checks import CATEGORIES, NEEDS_A_LOOK, included, needs_review
from .schemas import Receipt, Report

Lang = Literal["ko", "en"]

LABELS: dict[str, dict] = {
    "ko": {
        "title": "경비 정산서",
        "sheet": "정산서",
        "excluded_sheet": "제외",
        "period": "기간",
        "receipts": "영수증",
        "made": "작성",
        "count_unit": "장",
        "review_note": "노란 줄 {n}건은 확인이 필요해요.",
        "columns": ["번호", "날짜", "시간", "사용처", "사업자번호", "내용", "분류", "결제", "통화", "금액", "부가세",
                    "확인 필요", "사진 파일"],
        "excluded_columns": ["사진 파일", "이유"],
        "by_category": "분류별 합계",
        "total": "합계",
        "reviewed": "확인함",
        "categories": {"meals": "식대", "cafe": "카페·음료", "transport": "교통비", "lodging": "숙박비",
                       "supplies": "소모품비", "other": "기타"},
        "payments": {"card": "카드", "cash": "현금", "other": "기타", "unknown": "모름"},
        "fields": {"merchant": "사용처", "business_number": "사업자번호", "date": "날짜", "time": "시간",
                   "total": "금액", "breakdown": "금액 내역"},
        "flags": {
            "not_receipt": "영수증이 아니에요",
            "duplicate_photo": "같은 사진이 이미 있어요",
            "unreadable": "잘 안 보인 칸: {fields}",
            "no_total": "금액이 없어요",
            "bad_date": "날짜를 확인해 주세요",
            "sum_mismatch": "영수증 속 금액끼리 안 맞아요: {lines} = {sum}, 합계 {total}",
            "bad_business_number": "사업자번호 검증 숫자가 안 맞아요",
            "no_business_number": "사업자번호가 비어 있어요",
            "possible_duplicate": "같은 날짜·금액의 영수증이 이미 있어요",
        },
    },
    "en": {
        "title": "Expense report",
        "sheet": "Report",
        "excluded_sheet": "Left out",
        "period": "Period",
        "receipts": "Receipts",
        "made": "Made",
        "count_unit": "",
        "review_note": "{n} yellow rows need a look.",
        "columns": ["No.", "Date", "Time", "Merchant", "Business no.", "Description", "Category", "Payment",
                    "Currency", "Amount", "Tax", "Needs a look", "Photo"],
        "excluded_columns": ["Photo", "Reason"],
        "by_category": "Totals by category",
        "total": "Total",
        "reviewed": "Checked",
        "categories": {"meals": "Meals", "cafe": "Café", "transport": "Transport", "lodging": "Lodging",
                       "supplies": "Supplies", "other": "Other"},
        "payments": {"card": "Card", "cash": "Cash", "other": "Other", "unknown": "Unknown"},
        "fields": {"merchant": "merchant", "business_number": "business number", "date": "date", "time": "time",
                   "total": "total", "breakdown": "breakdown"},
        "flags": {
            "not_receipt": "Not a receipt",
            "duplicate_photo": "The same photo is already in",
            "unreadable": "Hard to read: {fields}",
            "no_total": "No total",
            "bad_date": "Check the date",
            "sum_mismatch": "The receipt's own numbers don't add up: {lines} = {sum}, total {total}",
            "bad_business_number": "Business number fails its check digit",
            "no_business_number": "No business number",
            "possible_duplicate": "Same date and amount as another receipt",
        },
    },
}
TAX_LABEL = re.compile(r"부가세|부가가치세|세액|\b(?:vat|gst|tax)\b", re.IGNORECASE)  # not "Taxi fare"
BEFORE_TAX = re.compile(r"공급|과세|면세|before|pre-?tax|excl|별도", re.IGNORECASE)  # "공급가액(부가세 별도)" is not tax
MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
FIRST_ROW = 6  # rows 1-4 hold the title block, row 5 the header

YELLOW = PatternFill("solid", fgColor="FFF4CE")
GREY = PatternFill("solid", fgColor="EEF1F5")
WIDTHS = [6, 12, 7, 24, 14, 24, 11, 8, 7, 13, 11, 34, 22]


def plain_number(amount: float, currency: str) -> str:
    return f"{amount:,.0f}" if currency in nz.ZERO_DECIMAL else f"{amount:,.2f}"


def flag_text(row: Receipt, flag: str, lang: Lang) -> str:
    t = LABELS[lang]
    fields = ", ".join(t["fields"][name] for name in row.unreadable)
    # A sum mismatch shows its arithmetic, so the person knows which numbers to compare with the photo.
    signs = [("-" if line.amount < 0 else "") if i == 0 else (" - " if line.amount < 0 else " + ")
             for i, line in enumerate(row.breakdown)]
    lines = "".join(sign + plain_number(abs(line.amount), row.currency) for sign, line in zip(signs, row.breakdown))
    added = nz.round_amount(sum(line.amount for line in row.breakdown), row.currency)
    total = plain_number(row.total, row.currency) if row.total is not None else "?"
    return t["flags"][flag].format(fields=fields, lines=lines, sum=plain_number(added, row.currency), total=total)


def check_text(row: Receipt, lang: Lang) -> str:
    if row.reviewed and NEEDS_A_LOOK.intersection(row.flags):
        return LABELS[lang]["reviewed"]
    return "; ".join(flag_text(row, flag, lang) for flag in row.flags if flag in NEEDS_A_LOOK)


def tax(row: Receipt) -> float | None:
    """The tax lines of the breakdown (부가세, VAT, Tax), for the column accountants ask for."""
    amounts = [
        line.amount for line in row.breakdown if TAX_LABEL.search(line.label) and not BEFORE_TAX.search(line.label)
    ]
    return nz.round_amount(sum(amounts), row.currency) if amounts else None


def report_rows(report: Report) -> list[Receipt]:
    """Rows that count, oldest first; rows without a readable date go last."""
    rows = [row for row in report.rows if included(row)]
    return sorted(rows, key=lambda row: (nz.normalize_date(row.date) or "9999", row.time, row.file_name))


def cells(row: Receipt, lang: Lang) -> list:
    t = LABELS[lang]
    return [
        row.date,
        row.time,
        row.merchant,
        row.business_number,
        row.summary,
        t["categories"][row.category],
        t["payments"][row.payment],
        row.currency,
        row.total,
        tax(row),
        check_text(row, lang),
        row.file_name,
    ]


def csv_safe(value):
    """Text that a spreadsheet would run as a formula gets a leading apostrophe, as OWASP advises for CSV files."""
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    return value


def put(ws, r: int, col: int, value):
    """Writes a cell. Text read off a receipt stays text: "=HYPERLINK(...)" must never become a formula."""
    cell = ws.cell(r, col, value)
    if isinstance(value, str):
        cell.data_type = "s"
    return cell


def money_format(currency: str) -> str:
    return "#,##0" if currency in nz.ZERO_DECIMAL else "#,##0.00"


def to_xlsx(report: Report, lang: Lang, today: date) -> bytes:
    t = LABELS[lang]
    rows = report_rows(report)
    summary = report.summary
    wb = Workbook()
    ws = wb.active
    ws.title = t["sheet"]

    ws["A1"] = t["title"]
    ws["A1"].font = Font(size=16, bold=True)
    period = " ~ ".join(summary.period) if summary.period else "-"
    ws["A2"] = f"{t['period']}: {period} · {t['receipts']}: {summary.count}{t['count_unit']} · {t['made']}: {today}"
    if summary.needs_review:
        ws["A3"] = t["review_note"].format(n=summary.needs_review)

    header_row = FIRST_ROW - 1
    thin = Side(style="thin", color="B8C0CC")
    for col, name in enumerate(t["columns"], 1):
        cell = ws.cell(header_row, col, name)
        cell.font = Font(bold=True)
        cell.fill = GREY
        cell.border = Border(bottom=thin)
        cell.alignment = Alignment(horizontal="center")

    last = FIRST_ROW + len(rows) - 1
    for i, row in enumerate(rows):
        r = FIRST_ROW + i
        values = [i + 1, *cells(row, lang)]
        iso = nz.normalize_date(row.date)
        values[1] = date.fromisoformat(iso) if iso else row.date  # a real date, so Excel can sort and filter
        for col, value in enumerate(values, 1):
            cell = put(ws, r, col, value)
            if needs_review(row):
                cell.fill = YELLOW
        ws.cell(r, 2).number_format = "yyyy-mm-dd"
        for col in (10, 11):
            ws.cell(r, col).number_format = money_format(row.currency)

    # Totals are formulas, so fixing an amount in Excel updates them.
    r = last + 2
    ws.cell(r, 4, t["by_category"]).font = Font(bold=True)
    if rows:
        amount, category, currency = (f"${c}${FIRST_ROW}:${c}${last}" for c in ("J", "G", "I"))
        for code in summary.totals:
            for key in CATEGORIES:
                if key in summary.totals[code].by_category:
                    r += 1
                    ws.cell(r, 4, t["categories"][key])
                    ws.cell(r, 9, code)
                    ws.cell(r, 10, f"=SUMIFS({amount},{category},D{r},{currency},I{r})")
                    ws.cell(r, 10).number_format = money_format(code)
            r += 1
            ws.cell(r, 4, f"{t['total']} ({code})").font = Font(bold=True)
            ws.cell(r, 9, code)
            ws.cell(r, 10, f"=SUMIFS({amount},{currency},I{r})")
            ws.cell(r, 10).number_format = money_format(code)
            ws.cell(r, 10).font = Font(bold=True)
            r += 1

    for col, width in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = f"A{FIRST_ROW}"
    if rows:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(WIDTHS))}{last}"

    left_out = [row for row in report.rows if not included(row)]
    if left_out:
        ex = wb.create_sheet(t["excluded_sheet"])
        ex.append(t["excluded_columns"])
        for cell in ex[1]:
            cell.font = Font(bold=True)
            cell.fill = GREY
        for i, row in enumerate(left_out, 2):
            put(ex, i, 1, row.file_name)
            put(ex, i, 2, "; ".join(flag_text(row, flag, lang) for flag in row.flags))
        ex.column_dimensions["A"].width = 26
        ex.column_dimensions["B"].width = 34

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def to_csv(report: Report, lang: Lang) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(LABELS[lang]["columns"])
    for i, row in enumerate(report_rows(report), 1):
        values = [csv_safe(value) for value in cells(row, lang)]
        for col in (8, 9):  # amount and tax: plain numbers, no thousands separators
            if values[col] is not None:
                values[col] = f"{values[col]:.0f}" if row.currency in nz.ZERO_DECIMAL else f"{values[col]:.2f}"
        writer.writerow([i, *values])
    # The byte order mark makes Excel open the Korean text as UTF-8 instead of guessing a legacy encoding.
    return ("﻿" + buf.getvalue()).encode("utf-8")
