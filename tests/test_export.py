import csv
import io
from datetime import datetime

from openpyxl import load_workbook

from app.checks import build_report
from app.export import flag_text, tax, to_csv, to_xlsx
from conftest import TODAY


def sample_report(make_row):
    return build_report(
        [
            make_row(id="1" * 16, file_name="b.jpg", date="2026-10-15", time="12:31", merchant="온누리김밥",
                     total=8500, breakdown=[], category="meals", business_number="214-03-55124"),
            make_row(id="2" * 16, file_name="a.jpg"),
            make_row(id="3" * 16, file_name="usd.jpg", currency="USD", business_number="", date="2026-10-13",
                     merchant="Blue Harbor Coffee", total=12.86,
                     breakdown=[{"label": "Subtotal", "amount": 10}, {"label": "Tax", "amount": 0.86},
                                {"label": "Tip", "amount": 2}]),
            make_row(id="4" * 16, file_name="odd.jpg", time="18:40", total=6000),  # breakdown says 9,000
            make_row(id="5" * 16, file_name="note.jpg", is_receipt=False, total=None, date="", merchant=""),
        ],
        TODAY,
    )


def test_xlsx_lists_rows_oldest_first_with_totals_as_formulas(make_row):
    wb = load_workbook(io.BytesIO(to_xlsx(sample_report(make_row), "ko", TODAY)))
    ws = wb["정산서"]

    assert ws["A1"].value == "경비 정산서"
    assert ws["A2"].value == "기간: 2026-10-13 ~ 2026-10-15 · 영수증: 4장 · 작성: 2026-10-20"
    assert ws["A3"].value == "노란 줄 1건은 확인이 필요해요."
    assert [c.value for c in ws[5]][:4] == ["번호", "날짜", "시간", "사용처"]
    assert [ws.cell(r, 13).value for r in range(6, 10)] == ["usd.jpg", "a.jpg", "odd.jpg", "b.jpg"]
    assert ws["B6"].value == datetime(2026, 10, 13)  # a real date cell
    assert ws["J6"].number_format == "#,##0.00" and ws["J7"].number_format == "#,##0"
    assert ws["K7"].value == 818  # 부가세 from the breakdown
    assert ws["K6"].value == 0.86
    assert ws["L8"].value == "영수증 속 금액끼리 안 맞아요: 8,182 + 818 = 9,000, 합계 6,000"
    assert ws["A8"].fill.fgColor.rgb.endswith("FFF4CE")
    assert ws["A7"].fill.fill_type is None

    totals = [[ws.cell(r, c).value for c in (4, 9, 10)] for r in range(11, ws.max_row + 1)]
    assert totals == [
        ["분류별 합계", None, None],
        ["식대", "KRW", "=SUMIFS($J$6:$J$9,$G$6:$G$9,D12,$I$6:$I$9,I12)"],
        ["카페·음료", "KRW", "=SUMIFS($J$6:$J$9,$G$6:$G$9,D13,$I$6:$I$9,I13)"],
        ["합계 (KRW)", "KRW", "=SUMIFS($J$6:$J$9,$I$6:$I$9,I14)"],
        [None, None, None],
        ["카페·음료", "USD", "=SUMIFS($J$6:$J$9,$G$6:$G$9,D16,$I$6:$I$9,I16)"],
        ["합계 (USD)", "USD", "=SUMIFS($J$6:$J$9,$I$6:$I$9,I17)"],
    ]
    assert ws.freeze_panes == "A6"

    left_out = wb["제외"]
    assert [c.value for c in left_out[2]] == ["note.jpg", "영수증이 아니에요"]


def test_xlsx_in_english(make_row):
    wb = load_workbook(io.BytesIO(to_xlsx(sample_report(make_row), "en", TODAY)))
    ws = wb["Report"]

    assert ws["A1"].value == "Expense report"
    assert ws["L8"].value == "The receipt's own numbers don't add up: 8,182 + 818 = 9,000, total 6,000"
    assert wb["Left out"]["B2"].value == "Not a receipt"


def test_xlsx_with_nothing_to_report(make_row):
    wb = load_workbook(io.BytesIO(to_xlsx(build_report([], TODAY), "ko", TODAY)))
    assert wb["정산서"]["A2"].value.startswith("기간: - ")


def test_csv_opens_in_excel_as_utf8_and_has_plain_numbers(make_row):
    data = to_csv(sample_report(make_row), "ko")

    assert data.startswith("﻿".encode())
    rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
    assert rows[0][:4] == ["번호", "날짜", "시간", "사용처"]
    assert [r[12] for r in rows[1:]] == ["usd.jpg", "a.jpg", "odd.jpg", "b.jpg"]
    assert rows[1][9:11] == ["12.86", "0.86"]
    assert rows[2][9:11] == ["9000", "818"]
    assert rows[4][10] == ""


def test_tax_column_reads_only_tax_lines(make_row):
    assert tax(make_row()) == 818
    assert tax(make_row(breakdown=[{"label": "공급가액(부가세 별도)", "amount": 8182}, {"label": "VAT", "amount": 818}])) == 818
    assert tax(make_row(currency="USD", breakdown=[{"label": "Taxi fare", "amount": 18}])) is None
    assert tax(make_row(breakdown=[])) is None


def test_sum_mismatch_shows_the_arithmetic(make_row):
    parking = make_row(total=7000, breakdown=[{"label": "주차요금", "amount": 12000}, {"label": "할인", "amount": 6000}])
    assert flag_text(parking, "sum_mismatch", "en") == (
        "The receipt's own numbers don't add up: 12,000 - 6,000 = 6,000, total 7,000"
    )
    usd = make_row(currency="USD", business_number="", total=12.68,
                   breakdown=[{"label": "Subtotal", "amount": 10}, {"label": "Tax", "amount": 0.86},
                              {"label": "Tip", "amount": 2}])
    assert flag_text(usd, "sum_mismatch", "ko") == "영수증 속 금액끼리 안 맞아요: 10.00 + 0.86 + 2.00 = 12.86, 합계 12.68"


def test_text_from_a_receipt_never_becomes_a_formula(make_row):
    trap = '=HYPERLINK("http://example.com","click")'
    report = build_report([make_row(merchant=trap, file_name="+cmd.jpg", total=-6000, breakdown=[])], TODAY)

    ws = load_workbook(io.BytesIO(to_xlsx(report, "ko", TODAY)))["정산서"]
    assert ws["D6"].value == trap and ws["D6"].data_type == "s"
    assert ws["M6"].data_type == "s"

    rows = list(csv.reader(io.StringIO(to_csv(report, "ko").decode("utf-8-sig"))))
    assert rows[1][3] == "'" + trap
    assert rows[1][12] == "'+cmd.jpg"
    assert rows[1][9] == "-6000"  # a negative amount is a number, not text to escape
